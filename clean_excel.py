import openpyxl
import os
import sys

print("Loading original workbook (this will take ~40 seconds)...")
sys.stdout.flush()
wb = openpyxl.load_workbook('Redwed ledger.xlsx')

print("Cleaning sheets...")
sys.stdout.flush()
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    max_row = ws.max_row
    # Find the actual last row with data
    real_max = 1
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(v is not None for v in row):
            real_max = row_idx
            
    print(f"Sheet '{sheet_name}': max_row is {max_row}, real_max is {real_max}")
    sys.stdout.flush()
    if real_max < max_row:
        # Delete rows from real_max+1 to max_row
        print(f"  Deleting {max_row - real_max} empty rows...")
        sys.stdout.flush()
        ws.delete_rows(real_max + 1, max_row - real_max)

print("Saving cleaned workbook...")
sys.stdout.flush()
wb.save('Redwed ledger.xlsx')
print("Done! Check new file size.")
sys.stdout.flush()
